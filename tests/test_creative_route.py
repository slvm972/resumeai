"""
Тесты подключения creative-режима к /api/improve (шаг 2б) и защиты /api/improve/docx.
Без сети: creative_improve и _run_improve_pipeline мокаются.
Запуск: FLASK_ENV=testing python -m pytest tests/test_creative_route.py -q
"""
import io
import logging
import os
import sys
import uuid
from unittest.mock import patch

import pytest

sys.path.insert(0, '/home/claude/resumeai')

os.environ['FLASK_ENV'] = 'testing'
os.environ.setdefault('SECRET_KEY', 'test-secret-key-for-tests-only')
os.environ.setdefault('JWT_SECRET_KEY', 'test-jwt-secret-for-tests-only')
os.environ.setdefault('GROQ_API_KEY', 'test-groq-key-not-used-mocked-out')

import importlib
import config as config_module
importlib.reload(config_module)

from app import create_app, db

_KEY = "SECRET-GROQ-KEY-ROUTE-TEST"
_TEXT = "Anna Smith\nanna@example.com\nSales Associate, Acme Store, 2018-2022\nWorked with customers"
_MARKER = "UNIQUE-RESUME-MARKER-4242"
_CHANGELOG = ("Changes\nRewritten:\n\u00abWorked with customers\u00bb \u2192 \u00abServed customers\u00bb\n"
              "Added:\n+ \u00abSummary\u00bb\nRemoved:\n- \u00abOld line\u00bb")

_PRECISE_RESULT = {
    'success': True,
    'improved_resume': '###ITEM_001###\nPrecise text',
    'display_text': 'Precise text',
    'detected_language': 'English',
    'item_ids': ['001'],
    'quality_report': {'summary': {'accepted': 1}, 'changes': []},
    'has_original_docx': False,
    'creativity_mode': 'precise',
    'guardrail_report': None,
    'guardrail_rejected': False,
}


def _creative_ok(**over):
    r = {"attempted": True, "acceptable": True, "status": "clean",
         "improved_resume": "Anna Smith\nanna@example.com\nSummary\nSales associate.\n- Served customers",
         "display_text": "Anna Smith\nanna@example.com\nSummary\nSales associate.\n- Served customers",
         "derived": "- 4 years", "changelog": _CHANGELOG, "detected_language": "English",
         "rounds": 0, "calls": 2, "prompt_version": "v2.2", "checks": {}, "issues": [], "error": None}
    r.update(over)
    return r


@pytest.fixture
def app_ctx():
    app = create_app('testing')
    app.config['TESTING'] = True
    app.config['CREATIVE_V2_ENABLED'] = 'true'
    app.config['GROQ_API_KEY'] = _KEY
    with app.app_context():
        yield app


def _login(app, client, admin=False):
    from app.services.auth_service import AuthService
    email = 'creative-route-%s@example.com' % uuid.uuid4().hex[:10]
    user = AuthService.register(email, 'somepassword123')['user']
    with client.session_transaction() as sess:
        sess['user_id'] = user.id
        if admin:
            sess['admin'] = 'admin'
    return user


def _json_post(client, mode='creative', text=_TEXT):
    return client.post('/api/improve', json={'resume_text': text, 'creativity_mode': mode})


def _file_post(client, name, content, mode='creative'):
    return client.post('/api/improve', data={'file': (io.BytesIO(content), name), 'creativity_mode': mode},
                       content_type='multipart/form-data')


def _docx_bytes(text="ORIGINAL TEXT FROM FILE"):
    from docx import Document
    d = Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _used(user):
    sub = user.get_active_subscription()
    db.session.refresh(sub)
    return sub.credits_used, sub.improvement_used


# ---- флаг и режим -------------------------------------------------------------

def test_01_flag_off_creative_uses_precise_pipeline_without_bridge(app_ctx):
    app_ctx.config['CREATIVE_V2_ENABLED'] = 'false'
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with patch('app.creative_bridge.creative_improve') as bridge, \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT) as pipe:
        resp = _json_post(client)
    assert resp.status_code == 200 and bridge.call_count == 0
    assert pipe.call_args[0][4] == 'creative'                       # исходное значение режима
    body = resp.get_json()
    assert 'engine' not in body and 'creative_fallback' not in body


def test_02_acceptable_creative_response_and_one_credit(app_ctx):
    client = app_ctx.test_client()
    user = _login(app_ctx, client)
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok()) as bridge, \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT) as pipe:
        resp = _json_post(client)
    assert resp.status_code == 200 and bridge.call_count == 1 and pipe.call_count == 0
    b = resp.get_json()
    assert b['success'] is True and b['engine'] == 'creative_v2' and b['creativity_mode'] == 'creative'
    assert b['item_ids'] == [] and b['has_original_docx'] is False
    assert b['guardrail_report'] is None and b['guardrail_rejected'] is False
    assert b['improved_resume'] == b['display_text'] and 'DERIVED' not in b['improved_resume']
    assert b['original_language'] == 'English' and b['original_format'] == 'txt'
    assert b['creative_status'] == 'clean' and b['creative_rounds'] == 0
    assert b['derived'] == '- 4 years' and b['changelog'].startswith('Changes')
    assert _used(user) == (1, 1)                                     # ровно один кредит


def test_03_quality_report_is_frontend_compatible(app_ctx):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok()):
        qr = _json_post(client).get_json()['quality_report']
    assert qr['summary']['accepted'] >= 1 and qr['engine'] == 'creative_v2'
    assert qr['changes'][0] == {'original_text': 'Worked with customers', 'improved_text': 'Served customers'}
    assert {'original_text': '', 'improved_text': 'Summary'} in qr['changes']
    assert {'original_text': 'Old line', 'improved_text': ''} in qr['changes']
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok(changelog='')):
        assert _json_post(client).get_json()['quality_report']['summary']['accepted'] == 1


def test_04_not_acceptable_falls_back_to_precise(app_ctx):
    client = app_ctx.test_client()
    user = _login(app_ctx, client)
    bad = _creative_ok(acceptable=False, status='unresolved')
    with patch('app.creative_bridge.creative_improve', return_value=bad), \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT) as pipe:
        resp = _json_post(client)
    assert resp.status_code == 200 and pipe.call_args[0][4] == 'precise'
    b = resp.get_json()
    assert b['engine'] == 'precise' and b['creative_fallback'] is True and b['creative_status'] == 'unresolved'
    assert b['item_ids'] == ['001']
    assert _used(user) == (1, 1)                                     # кредит по правилам точного пути


def test_05_bridge_exception_is_not_500_and_leaks_nothing(app_ctx):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with patch('app.creative_bridge.creative_improve', side_effect=RuntimeError("secret-detail " + _MARKER)), \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT) as pipe:
        resp = _json_post(client)
    assert resp.status_code == 200 and pipe.call_args[0][4] == 'precise'
    b = resp.get_json()
    assert b['engine'] == 'precise' and b['creative_fallback'] is True and b['creative_status'] == 'error'
    assert b'secret-detail' not in resp.get_data() and _MARKER.encode() not in resp.get_data()


def test_06_precise_mode_with_flag_on_does_not_call_bridge(app_ctx):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with patch('app.creative_bridge.creative_improve') as bridge, \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT) as pipe:
        resp = _json_post(client, mode='precise')
    assert resp.status_code == 200 and bridge.call_count == 0 and pipe.call_args[0][4] == 'precise'


@pytest.mark.parametrize("name", ["cv.pdf", "cv.odt", "cv.rtf", "cv.doc", "cv.txt"])
def test_07_non_docx_files_skip_creative(app_ctx, name):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with patch('app.creative_bridge.creative_improve') as bridge, \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT) as pipe:
        resp = _file_post(client, name, b'resume text ' * 10)
    assert resp.status_code == 200 and bridge.call_count == 0 and pipe.call_count == 1


def test_08_docx_file_goes_to_bridge_and_sets_no_session_tokens(app_ctx):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with client.session_transaction() as sess:                       # «старые» значения от прежнего запроса
        sess['original_docx_token'] = 'old-token'
        sess['item_ids'] = ['001', '002']
    data = _docx_bytes()
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok()) as bridge, \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT) as pipe:
        resp = _file_post(client, 'cv.docx', data)
    assert resp.status_code == 200 and pipe.call_count == 0
    assert bridge.call_args[0][0] == data and bridge.call_args[0][1] == 'cv.docx'
    b = resp.get_json()
    assert b['engine'] == 'creative_v2' and b['original_format'] == 'docx' and b['has_original_docx'] is False
    with client.session_transaction() as sess:
        assert 'original_docx_token' not in sess and 'item_ids' not in sess


# ---- кредиты и доступ ---------------------------------------------------------

def test_10_admin_session_is_not_charged(app_ctx):
    client = app_ctx.test_client()
    user = _login(app_ctx, client, admin=True)
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok()):
        resp = _json_post(client)
    assert resp.status_code == 200 and resp.get_json()['engine'] == 'creative_v2'
    assert _used(user) == (0, 0)


def test_11_no_credits_gives_403_before_bridge(app_ctx):
    client = app_ctx.test_client()
    user = _login(app_ctx, client)
    sub = user.get_active_subscription()
    sub.credits_used = sub.credits_granted
    db.session.commit()
    with patch('app.creative_bridge.creative_improve') as bridge:
        resp = _json_post(client)
    assert resp.status_code == 403 and bridge.call_count == 0
    assert _used(user)[0] == sub.credits_granted


def test_12_not_acceptable_creative_is_not_charged_by_itself(app_ctx):
    client = app_ctx.test_client()
    user = _login(app_ctx, client)
    no_gain = dict(_PRECISE_RESULT, quality_report={'summary': {'accepted': 0}, 'changes': []})
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok(acceptable=False)), \
         patch('app.missing_routes4._run_improve_pipeline', return_value=no_gain):
        assert _json_post(client).status_code == 200
    assert _used(user) == (0, 0)                                     # списание только по правилам точного пути


def test_13_attempted_false_keeps_previous_behaviour(app_ctx):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    short = {"attempted": False, "acceptable": False, "error": "Resume text too short"}
    with patch('app.creative_bridge.creative_improve', return_value=short), \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT) as pipe:
        resp = _json_post(client)
    assert pipe.call_args[0][4] == 'creative' and 'creative_fallback' not in resp.get_json()


# ---- утечки и логи --------------------------------------------------------------

def test_20_groq_key_not_in_any_response(app_ctx):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok()):
        assert _KEY.encode() not in _json_post(client).get_data()
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok(acceptable=False)), \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT):
        assert _KEY.encode() not in _json_post(client).get_data()
    with patch('app.creative_bridge.creative_improve', side_effect=RuntimeError(_KEY)):
        with patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT):
            assert _KEY.encode() not in _json_post(client).get_data()


def test_21_log_line_has_engine_status_rounds_calls_and_no_resume_text(app_ctx, caplog):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    caplog.set_level(logging.INFO)
    issue = [{"type": "FLUFF", "quote": _MARKER}]
    with patch('app.creative_bridge.creative_improve', return_value=_creative_ok(issues=issue)):
        _json_post(client, text=_TEXT + " " + _MARKER)
    with patch('app.creative_bridge.creative_improve', side_effect=RuntimeError("detail " + _MARKER)), \
         patch('app.missing_routes4._run_improve_pipeline', return_value=_PRECISE_RESULT):
        _json_post(client, text=_TEXT + " " + _MARKER)
    assert "engine=creative_v2 status=clean rounds=0 calls=2" in caplog.text
    assert "engine=precise creative_fallback=1 status=error" in caplog.text
    assert _MARKER not in caplog.text and _KEY not in caplog.text and "detail" not in caplog.text


# ---- защита /api/improve/docx ---------------------------------------------------

def _post_docx(client, text, original=None):
    data = {'improved_resume': text}
    if original is not None:
        data['original_file'] = (io.BytesIO(original), 'cv.docx')
    return client.post('/api/improve/docx', data=data, content_type='multipart/form-data')


def test_30_docx_text_without_markers_builds_plain_docx_not_original(app_ctx):
    from docx import Document
    client = app_ctx.test_client()
    _login(app_ctx, client)
    text = "Anna Smith\nanna@example.com\nSummary\nSales associate.\n- Served customers"
    with patch('app.missing_routes4._apply_improved_text_to_docx') as apply_:
        resp = _post_docx(client, text, original=_docx_bytes())
    assert resp.status_code == 200 and apply_.call_count == 0
    paras = [p.text for p in Document(io.BytesIO(resp.get_data())).paragraphs]
    assert paras == ["Anna Smith", "anna@example.com", "Summary", "Sales associate.", "Served customers"] \
        or paras == ["Anna Smith", "anna@example.com", "Summary", "Sales associate.", "- Served customers"]
    assert "ORIGINAL TEXT FROM FILE" not in paras


def test_31_docx_text_with_markers_still_applies_to_original(app_ctx):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with patch('app.missing_routes4._apply_improved_text_to_docx', return_value=io.BytesIO(b'APPLIED')) as apply_:
        resp = _post_docx(client, "###ITEM_001###\nNew text", original=_docx_bytes())
    assert resp.status_code == 200 and apply_.call_count == 1 and resp.get_data() == b'APPLIED'


def test_32_docx_without_markers_and_without_original_file_still_works(app_ctx):
    from docx import Document
    client = app_ctx.test_client()
    _login(app_ctx, client)
    resp = _post_docx(client, "Line one\nLine two")
    assert resp.status_code == 200
    assert [p.text for p in Document(io.BytesIO(resp.get_data())).paragraphs] == ["Line one", "Line two"]


def test_33_docx_export_clears_session_tokens(app_ctx):
    client = app_ctx.test_client()
    _login(app_ctx, client)
    with client.session_transaction() as sess:
        sess['original_docx_token'] = 'tok'
        sess['item_ids'] = ['001']
    _post_docx(client, "Line one\nLine two")
    with client.session_transaction() as sess:
        assert 'original_docx_token' not in sess and 'item_ids' not in sess
