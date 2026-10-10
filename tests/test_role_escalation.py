"""
Тесты для Cycle R1 — Role Escalation Guard: промпт-правило 14 +
отдельная детерминированная validation-функция _check_role_escalation.

Запуск: python -m pytest tests/test_role_escalation.py -v
"""
import io
import re
import sys
import pytest
from unittest.mock import patch, MagicMock
from docx import Document

import app.missing_routes4 as mr


# ===========================================================================
# R1_1 — промпты содержат Rule 14
# ===========================================================================

def test_R1_1_standard_prompt_has_rule14():
    prompt = mr._build_standard_system_prompt(3, "Russian")
    assert "STRICT ROLE PRESERVATION" in prompt
    assert "руководил" in prompt and "led" in prompt


def test_R1_1_relaxed_prompt_has_rule14():
    prompt = mr._build_plain_relaxed_system_prompt(1, "Russian")
    assert "STRICT ROLE PRESERVATION" in prompt


def test_R1_1_summary_prompt_has_rule14():
    prompt = mr._build_summary_repositioning_prompt(1, "Russian")
    assert "STRICT ROLE PRESERVATION" in prompt


# ===========================================================================
# R1_2 — старые конфликтующие инструкции убраны
# ===========================================================================

def test_R1_2_relaxed_prompt_no_longer_encourages_led_escalation():
    prompt_lower = mr._build_plain_relaxed_system_prompt(1, "Russian").lower()
    assert 'was responsible for" -> "led' not in prompt_lower
    assert '"managed" -> "directed"' not in prompt_lower
    assert '-> "led"' not in prompt_lower
    assert '-> "directed"' not in prompt_lower
    # Позитивная проверка — фикс реально на месте, не просто "плохого нет"
    assert '"helped with" -> "supported"' in prompt_lower
    assert '"worked on" -> "implemented"' in prompt_lower


def test_R1_2_managed_directed_example_removed_from_both_prompts():
    for prompt in (
        mr._build_standard_system_prompt(3, "Russian"),
        mr._build_plain_relaxed_system_prompt(1, "Russian"),
    ):
        assert "\"Managed\" → \"Directed\"" not in prompt


# ===========================================================================
# R1_3 — матрица эскалации: IC-формулировка НЕ должна пропускать
# management-маркер, появившийся только в новом тексте
# ===========================================================================

_ESCALATION_MATRIX = [
    ("Работал над проектом обновления Cisco", "Руководил проектом обновления Cisco"),
    ("Отвечал за сеть компании", "Руководил сетью компании"),
    ("Занимался технической поддержкой пользователей", "Возглавлял техническую поддержку пользователей"),
    ("Участвовал в проекте модернизации", "Курировал проект модернизации"),
    ("Поддерживал серверную инфраструктуру", "Управлял командой серверной инфраструктуры"),
    ("Выполнял задачи по обновлению оборудования", "Руководил обновлением оборудования"),
]


@pytest.mark.parametrize("orig,escalated", _ESCALATION_MATRIX)
def test_R1_3_escalation_matrix_rejected(orig, escalated):
    ok, reason = mr._check_role_escalation(orig, escalated)
    assert not ok, f"Эскалация не поймана: {orig!r} -> {escalated!r}"
    assert reason


# ===========================================================================
# R1_4 — genuine management/leadership должен сохраняться (guard не
# вмешивается, если у оригинала УЖЕ есть management-маркер)
# ===========================================================================

_MANAGEMENT_PRESERVATION_MATRIX = [
    ("Руководил командой из 5 инженеров", "Возглавлял команду из 5 инженеров"),
    ("Управлял отделом технической поддержки", "Руководил отделом технической поддержки"),
    ("Курировал работу команды из 8 специалистов", "Управлял командой из 8 специалистов"),
    ("Managed a team of 5 designers", "Directed a team of 5 designers"),
]


@pytest.mark.parametrize("orig,reworded", _MANAGEMENT_PRESERVATION_MATRIX)
def test_R1_4_genuine_management_preserved(orig, reworded):
    ok, reason = mr._check_role_escalation(orig, reworded)
    assert ok, f"Genuine management ошибочно отклонён: {orig!r} -> {reworded!r} ({reason})"


# ===========================================================================
# R1_5 — не про добавление новой ТЕХНИЧЕСКОЙ ответственности (не в
# management-регистре вообще, ни в orig, ни в new) — guard корректно НЕ
# трогает это (не его зона: это гипотетически ловится Fact Validation/
# промптом, не keyword-based role-escalation guard'ом; отмечаем как
# известное ограничение, не пытаемся регэкспом поймать открытый список
# "любое возможное техническое действие").
# ===========================================================================

def test_R1_5_non_management_technical_inflation_not_caught_by_this_guard():
    """Известное ограничение: "Отвечал за сеть" -> "Администрировал сеть"
    не содержит ни одного management/leadership-маркера ни до, ни после —
    _check_role_escalation его не ловит (это не его задача, а зона
    промпт-инструкции/Fact Validation). Тест документирует границу
    guard'а, а не требует от него невозможного."""
    ok, reason = mr._check_role_escalation("Отвечал за сеть", "Администрировал сеть")
    assert ok  # guard пропускает — вне его зоны ответственности по дизайну


def test_R1_5_role_escalation_guard_is_separate_from_fact_validation():
    """_validate_block НЕ трогали в этом цикле — её поведение на
    Managed->Directed не изменилось (regression guard)."""
    valid, _ = mr._validate_block("Managed a team", "Directed a team")
    assert valid


# ===========================================================================
# Интеграционный тест — guard реально подключён в pipeline
# ===========================================================================

# ===========================================================================
# R1.1-fix — LED (аппаратный компонент) и managed/directed в техническом
# контексте (managed switch/network, directed graph) НЕ должны ложно
# срабатывать как role escalation. Найдено пользователем после первого
# прохода Cycle R1 — reg-guard, чтобы не откатилось.
# ===========================================================================

_TECH_FALSE_POSITIVE_MATRIX = [
    ("Configured network switches", "Installed LED indicators on network switches"),
    ("Set up server racks", "Configured LED display panel for status monitoring"),
    ("Set up network equipment", "Deployed managed network infrastructure"),
    ("Configured basic switches", "Deployed managed switches across the office"),
    ("Configured basic firewall rules", "Managed firewall infrastructure and network devices"),
    ("Studied graph algorithms", "Implemented directed graph traversal"),
]


@pytest.mark.parametrize("orig,new", _TECH_FALSE_POSITIVE_MATRIX)
def test_R1_6_technical_usage_not_flagged_as_escalation(orig, new):
    ok, reason = mr._check_role_escalation(orig, new)
    assert ok, f"Ложное срабатывание на техническом употреблении: {orig!r} -> {new!r} ({reason})"


def test_R1_6_led_still_catches_genuine_verb_usage():
    """Regression guard в обратную сторону: исключение LED-аббревиатуры не
    должно съедать реальный глагол "led"/"Led" (обычный регистр)."""
    ok, reason = mr._check_role_escalation(
        "Worked on the migration project", "Led the migration project"
    )
    assert not ok, "Настоящий глагол 'led' должен по-прежнему ловиться"


def test_R1_6_managed_still_catches_genuine_people_management():
    ok, reason = mr._check_role_escalation(
        "Worked on the design team", "Managed the design team"
    )
    assert not ok, "Настоящее 'managed [людей]' должно по-прежнему ловиться"


def test_R1_integration_escalation_rejected_in_full_pipeline():
    doc = Document()
    doc.add_paragraph("John Smith")
    doc.add_paragraph("john.smith@example.com")
    doc.add_paragraph("Отвечал за сеть компании из 50 станций", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    def fake_post(url, headers=None, json=None, timeout=None):
        content = "###ITEM_003###\nРуководил сетью компании из 50 станций"
        resp = MagicMock()
        resp.status_code = 200
        resp.json.return_value = {
            "choices": [{"message": {"content": content}}],
            "usage": {"total_tokens": 10},
        }
        return resp

    with patch("requests.post", side_effect=fake_post):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key", creativity_mode="precise",
        )

    assert result["success"], result.get("error")
    assert "Руководил" not in result["display_text"]
    assert "Отвечал за сеть" in result["display_text"]
    assert result["quality_report"]["summary"]["rejected_role_escalation"] == 1


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
