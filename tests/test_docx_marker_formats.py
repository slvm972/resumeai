"""
Тесты форматов меток ###ITEM_NNN### / ###ITEM_NNN:TYPE### в экспорте DOCX
(фикс: /api/improve/docx и /api/admin/improve/docx не распознавали метки с типом).
Без сети: requests.post замокан, текст с метками генерирует настоящий
_run_improve_pipeline (а не придуманная строка).
Запуск: FLASK_ENV=testing python -m pytest tests/test_docx_marker_formats.py -q
"""
import io
import os
import re
import sys
from unittest.mock import patch

import pytest

sys.path.insert(0, '/home/claude/resumeai')

os.environ['FLASK_ENV'] = 'testing'
os.environ.setdefault('SECRET_KEY', 'test-secret-key-for-tests-only')
os.environ.setdefault('JWT_SECRET_KEY', 'test-jwt-secret-for-tests-only')
os.environ.setdefault('GROQ_API_KEY', 'test-groq-key-not-used-mocked-out')

from docx import Document

import app.missing_routes4 as mr
from app import create_app, _ITEM_MARKER_PATTERN

_P1 = "Professional Experience"
_P2 = "Managed a team of 5 designers across 3 product verticals"
_P3 = "Built internal tools for the sales department"
_NEW2 = "Directed a team of 5 designers across 3 product verticals"
_NEW3 = "Delivered internal tools for the sales department"


# ---------------------------------------------------------------------------
# Вспомогательное
# ---------------------------------------------------------------------------

def _docx_bytes():
    """Оригинал из 3 абзацев: заголовок (Heading 1) и два обычных."""
    d = Document()
    d.add_heading(_P1, level=1)
    d.add_paragraph(_P2)
    d.add_paragraph(_P3)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class _Resp:
    status_code = 200

    def __init__(self, content):
        self._c = content

    def json(self):
        return {"choices": [{"message": {"content": self._c}}], "usage": {"total_tokens": 5}}


def _echo_post(url, headers=None, json=None, timeout=None):
    """Мок Groq: возвращает блоки из запроса без изменений (формат ответа как у модели)."""
    user = json["messages"][-1]["content"]
    blocks = re.findall(r"###ITEM_\d+###\n.*?(?=\n\n###ITEM_|\n\nOUTPUT|\Z)", user, re.S)
    return _Resp("\n\n".join(blocks))


@pytest.fixture(scope="module")
def pipeline_result():
    """Результат настоящего пайплайна на docx из 3 абзацев (источник реального формата меток)."""
    with patch("requests.post", side_effect=_echo_post):
        res = mr._run_improve_pipeline(_docx_bytes(), "cv.docx", None, "fake-key")
    assert res["success"], res
    return res


@pytest.fixture
def typed_text(pipeline_result):
    """Текст с typed-метками от пайплайна, где абзацы 2 и 3 заменены на 'улучшенные'."""
    text = pipeline_result["improved_resume"]
    assert _P2 in text and _P3 in text
    return text.replace(_P2, _NEW2).replace(_P3, _NEW3)


@pytest.fixture
def untyped_text(typed_text):
    return re.sub(r"(###ITEM_\d+):\w+###", r"\1###", typed_text)


@pytest.fixture
def client():
    app = create_app('testing')
    app.config['TESTING'] = True
    with app.app_context():
        c = app.test_client()
        with c.session_transaction() as sess:
            sess['admin'] = 'admin'
        yield c


def _paras(buf_bytes):
    return [p.text for p in Document(io.BytesIO(buf_bytes)).paragraphs]


def _post(client, url, text, original=None, item_ids=None):
    data = {'improved_resume': text}
    if original is not None:
        data['original_file'] = (io.BytesIO(original), 'cv.docx')
    if item_ids is not None:
        import json
        data['item_ids'] = json.dumps(item_ids)
    return client.post(url, data=data, content_type='multipart/form-data')


def _assert_no_marker_leak(paras):
    for p in paras:
        assert "ITEM_" not in p and "###" not in p, repr(p)


_URLS = ['/api/improve/docx', '/api/admin/improve/docx']


# ---------------------------------------------------------------------------
# 1. Формат меток пайплайна и регулярка фикса
# ---------------------------------------------------------------------------

def test_pipeline_emits_typed_markers_recognized_by_fix_regex(pipeline_result):
    text = pipeline_result["improved_resume"]
    marker_lines = re.findall(r"^###ITEM_.*###$", text, re.M)
    assert len(marker_lines) == 3
    for line in marker_lines:
        assert re.fullmatch(r"###ITEM_\d{3}:(HEADING|BULLET|PLAIN|TABLE)###", line), line
    assert marker_lines[0].endswith(":HEADING###")          # Heading 1 -> HEADING
    assert len(re.findall(_ITEM_MARKER_PATTERN, text)) == 3  # регулярка фикса видит все
    assert not re.search(r"###ITEM_\d+###", text)            # старая регулярка typed не видела


def test_untyped_markers_recognized_by_fix_regex():
    text = mr._make_sep().format("001") + "\nText\n\n" + mr._make_sep().format("002") + "\nMore"
    assert text.startswith("###ITEM_001###")
    assert len(re.findall(_ITEM_MARKER_PATTERN, text)) == 2


# ---------------------------------------------------------------------------
# 2. Метки с типом + оригинальный файл: результат применён к оригиналу
# ---------------------------------------------------------------------------

@pytest.mark.parametrize('url', _URLS)
def test_typed_markers_with_original_applied_to_original(client, pipeline_result, typed_text, url):
    with patch('app.missing_routes4._apply_improved_text_to_docx', wraps=mr._apply_improved_text_to_docx) as spy:
        resp = _post(client, url, typed_text, original=_docx_bytes(), item_ids=pipeline_result["item_ids"])
    assert resp.status_code == 200
    assert spy.call_count == 1
    paras = _paras(resp.data)
    assert len(paras) == 3                                   # число абзацев то же
    assert paras == [_P1, _NEW2, _NEW3]                      # заменены абзацы 2 и 3
    _assert_no_marker_leak(paras)


# ---------------------------------------------------------------------------
# 3. Запасной путь без оригинала и токена
# ---------------------------------------------------------------------------

@pytest.mark.parametrize('url', _URLS)
def test_typed_markers_without_original_no_leak(client, typed_text, url):
    with patch('app.missing_routes4._apply_improved_text_to_docx') as spy:
        resp = _post(client, url, typed_text)
    assert resp.status_code == 200
    assert spy.call_count == 0
    paras = _paras(resp.data)
    _assert_no_marker_leak(paras)
    joined = "\n".join(paras)
    assert _P1 in joined and _NEW2 in joined and _NEW3 in joined


# ---------------------------------------------------------------------------
# 4. Метки без типа: прежнее поведение
# ---------------------------------------------------------------------------

@pytest.mark.parametrize('url', _URLS)
def test_untyped_markers_with_original_unchanged_behavior(client, pipeline_result, untyped_text, url):
    with patch('app.missing_routes4._apply_improved_text_to_docx', wraps=mr._apply_improved_text_to_docx) as spy:
        resp = _post(client, url, untyped_text, original=_docx_bytes(), item_ids=pipeline_result["item_ids"])
    assert resp.status_code == 200
    assert spy.call_count == 1
    assert _paras(resp.data) == [_P1, _NEW2, _NEW3]


@pytest.mark.parametrize('url', _URLS)
def test_untyped_markers_without_original_no_leak(client, untyped_text, url):
    resp = _post(client, url, untyped_text)
    assert resp.status_code == 200
    _assert_no_marker_leak(_paras(resp.data))


# ---------------------------------------------------------------------------
# 5. Текст без меток: простой DOCX (как раньше)
# ---------------------------------------------------------------------------

def test_text_without_markers_builds_plain_docx(client):
    text = "Line one\nLine two\nLine three"
    with patch('app.missing_routes4._apply_improved_text_to_docx') as spy:
        resp = _post(client, '/api/improve/docx', text, original=_docx_bytes())
    assert resp.status_code == 200
    assert spy.call_count == 0
    assert _paras(resp.data) == ["Line one", "Line two", "Line three"]


# ---------------------------------------------------------------------------
# 6. Экспорт PDF / ODT / RTF из текста с typed-метками: утечка меток?
#    (только проверка, исправлений нет)
# ---------------------------------------------------------------------------

def _pdf_text(buf):
    import PyPDF2
    return "\n".join(p.extract_text() or "" for p in PyPDF2.PdfReader(buf).pages)


def _odt_text(buf):
    from odf import teletype
    from odf.opendocument import load
    from odf.text import P
    doc = load(buf)
    return "\n".join(teletype.extractText(p) for p in doc.getElementsByType(P))


def _rtf_text(data):
    from striprtf.striprtf import rtf_to_text
    return rtf_to_text(data.decode('utf-8'))


@pytest.mark.parametrize('name, gen, extract', [
    ('pdf', mr._generate_pdf, _pdf_text),
    ('odt', mr._generate_odt, _odt_text),
    ('rtf', mr._generate_rtf, _rtf_text),
])
def test_export_formats_do_not_leak_typed_markers(typed_text, name, gen, extract):
    out = extract(gen(typed_text))
    assert "ITEM_" not in out and "###" not in out, (name, out[:200])
    assert _NEW2 in out.replace("\n", " ")
