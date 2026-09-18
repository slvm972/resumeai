"""
Тесты для Cycle B1 — Rule 12 в _build_standard_system_prompt() против
канцелярита и неестественных формулировок в буллетах Experience.

Bullet-элементы (как и heading/table) идут через _build_standard_system_prompt
(НЕ через _build_improve_prompt/_build_relaxed_prompt — таких функций в
кодовой базе нет, это неточность в исходном ТЗ Cycle B1). Правило 12
применяется в ОБОИХ режимах (precise и creative), т.к. bullet-группа не
имеет отдельного relaxed-варианта промпта по режиму — в отличие от
summary (Cycle S2, только creative) и plain (свой relaxed-промпт всегда).

Запуск: python -m pytest tests/test_experience_bullets.py -v
"""
import sys
import pytest

import app.missing_routes4 as mr


# ===========================================================================
# B1_1 — промпт для буллетов содержит явный запрет на канцелярит
# ===========================================================================

def test_B1_1_bullet_prompt_forbids_bureaucratic_verbs():
    prompt = mr._build_standard_system_prompt(3, "Russian")

    for forbidden in ("осуществлял", "выполнял задачи", "принимал участие", "занимался"):
        assert forbidden in prompt, \
            f"Промпт должен явно упоминать запрещённую форму {forbidden!r} как пример"

    assert "STRICTLY FORBIDDEN" in prompt


def test_B1_1_bullet_prompt_requires_strong_active_verbs():
    prompt = mr._build_standard_system_prompt(3, "Russian")
    for good_example in ("Внедрил", "Настроил", "Оптимизировал", "Провёл", "Сократил"):
        assert good_example in prompt


# ===========================================================================
# B1_2 — промпт содержит корректную замену для типового неестественного
# сочетания "внёс в эксплуатацию" (интерпретация: проверка на уровне
# текста промпта, а не новой fact-validation-логики — задача этого цикла
# ограничена промптом, _validate_block/_FABRICATED_CLAIM_RE не трогались
# и не рассчитаны на проверку идиоматичности глагольных сочетаний).
# ===========================================================================

def test_B1_2_commissioning_phrase_correction_present():
    prompt = mr._build_standard_system_prompt(3, "Russian")
    assert "внёс в эксплуатацию" in prompt, \
        "Промпт должен приводить некорректное сочетание как пример того, что запрещено"
    assert "ввёл в эксплуатацию" in prompt or "внедрил" in prompt.lower() or "Внедрил" in prompt, \
        "Промпт должен явно предлагать корректную замену"


# ===========================================================================
# Regression guard — правило 12 не должно ломать/дублироваться в других
# промптах (plain-relaxed и summary-repositioning в этом цикле НЕ
# трогались; у summary своё отдельное правило 12 с прошлого цикла — это
# ДРУГОЙ текст, не должен совпадать с bullet-версией дословно).
# ===========================================================================

def test_B1_regression_plain_relaxed_prompt_unaffected():
    prompt = mr._build_plain_relaxed_system_prompt(1, "Russian")
    assert "STRICTLY FORBIDDEN" not in prompt, \
        "Rule 12 из этого цикла адресован только _build_standard_system_prompt (bullet/heading/table)"


def test_B1_regression_summary_prompt_keeps_its_own_rule12():
    """summary-промпт получил СВОЁ правило 12 в Cycle S2 — другое по
    тексту (REPOSITION-контекст), но с той же общей идеей. Проверяем,
    что оно не перезаписано и не задвоено новым bullet-специфичным
    текстом этого цикла."""
    prompt = mr._build_summary_repositioning_prompt(1, "Russian")
    assert "REPOSITION" in prompt  # Cycle S2 правило 5 на месте
    assert "выполнял разнообразные задачи" in prompt  # Cycle S2 правило 12 на месте
    assert "STRICTLY FORBIDDEN" not in prompt, \
        "Bullet-специфичная формулировка Cycle B1 не должна была попасть в summary-промпт"


# ===========================================================================
# Cycle B2 — Few-Shot примеры в промпте + детерминированный Sanitizer
# ===========================================================================

def test_B2_1_bullet_prompt_contains_examples_block():
    prompt = mr._build_standard_system_prompt(3, "Russian")
    assert "EXAMPLES:" in prompt
    assert "Занимался внедрением системы" in prompt
    assert "Внёс в эксплуатацию систему" in prompt
    assert "Выполнял задачи по обновлению" in prompt
    # Каждый Bad-пример должен сопровождаться Good-вариантом рядом
    assert prompt.count("Bad:") == 3
    assert prompt.count("Good:") == 3


def test_B2_2_sanitizer_fixes_commissioning_phrase():
    result = mr._sanitize_awkward_phrasing("Внёс в эксплуатацию систему резервного копирования.")
    # Конкретно проверяем реально выбранную (см. _SANITIZER_REPLACEMENTS)
    # замену "ввёл в эксплуатацию" — не "внедрил" (второй вариант из ТЗ
    # структурно требует убрать "в эксплуатацию" из фразы целиком, что для
    # прямой regex-подстановки менее безопасно — см. докстринг функции).
    assert result == "Ввёл в эксплуатацию систему резервного копирования."


def test_B2_2_sanitizer_fixes_migration_phrase():
    result = mr._sanitize_awkward_phrasing("Осуществлял миграцию почтовой системы на Microsoft 365.")
    assert result == "Провёл миграцию почтовой системы на Microsoft 365."


def test_B2_2_sanitizer_case_preserving():
    # Строчная буква в середине предложения -> замена тоже строчная
    result = mr._sanitize_awkward_phrasing("Также внёс в эксплуатацию новую CRM.")
    assert result == "Также ввёл в эксплуатацию новую CRM."


def test_B2_2_sanitizer_handles_feminine_and_plural_forms():
    assert mr._sanitize_awkward_phrasing("Внесла в эксплуатацию новую CRM.") == \
        "Ввела в эксплуатацию новую CRM."
    assert mr._sanitize_awkward_phrasing("Внесли в эксплуатацию новую CRM.") == \
        "Ввели в эксплуатацию новую CRM."
    assert mr._sanitize_awkward_phrasing("Осуществляла миграцию базы данных.") == \
        "Провела миграцию базы данных."
    assert mr._sanitize_awkward_phrasing("Осуществляли миграцию базы данных.") == \
        "Провели миграцию базы данных."


def test_B2_2_sanitizer_noop_on_clean_text():
    clean = "Настроил VPN для 30 удалённых сотрудников и сократил время простоя на 40%."
    assert mr._sanitize_awkward_phrasing(clean) == clean


def test_B2_2_sanitizer_empty_and_none_safe():
    assert mr._sanitize_awkward_phrasing("") == ""
    assert mr._sanitize_awkward_phrasing(None) is None


def test_B2_regression_sanitizer_applied_before_validation_in_pipeline():
    """Интеграционная проверка: sanitizer реально подключён в
    _restore_and_validate (не только существует как отдельная функция) —
    LLM-ответ с "внёс в эксплуатацию" должен прийти в display_text уже
    санированным."""
    import io
    import re as re_module
    from unittest.mock import patch, MagicMock
    from docx import Document

    doc = Document()
    doc.add_paragraph("John Smith")
    doc.add_paragraph("john.smith@example.com")
    doc.add_paragraph(
        "Managed backup infrastructure and внёс в эксплуатацию new disaster recovery system",
        style="List Bullet",
    )
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    def _fake_post(url, headers=None, json=None, timeout=None):
        user_prompt = json["messages"][1]["content"]
        out_parts = []
        for m in re_module.finditer(
            r"###ITEM_(\d+)###\n(.*?)(?=\n\n###ITEM_|\n\nOUTPUT|\Z)",
            user_prompt, re_module.DOTALL,
        ):
            iid, text = m.group(1), m.group(2).strip()
            # LLM "игнорирует" промпт и всё равно генерирует плохую фразу —
            # ровно диагноз Cycle B2. Симулируем реалистичный ответ: замена
            # ведущего глагола синонимом (genuine word change, остаётся на
            # позиции 0 — не задевает Fact Validation по капитализации),
            # "внёс в эксплуатацию" оставляем нетронутым — его обязан
            # поймать sanitizer, не LLM.
            rewritten = text.replace("Managed", "Directed", 1)
            out_parts.append(f"###ITEM_{iid}###\n{rewritten}")
        content = "\n\n".join(out_parts)
        resp = MagicMock()
        resp.status_code = 200
        resp.json.return_value = {
            "choices": [{"message": {"content": content}}],
            "usage": {"total_tokens": 10},
        }
        return resp

    with patch("requests.post", side_effect=_fake_post):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="precise",
        )

    assert result["success"], result.get("error")
    assert "внёс в эксплуатацию" not in result["display_text"]
    assert "Внёс в эксплуатацию" not in result["display_text"]
    assert "ввёл в эксплуатацию" in result["display_text"] or \
           "Ввёл в эксплуатацию" in result["display_text"] or \
           "ввели в эксплуатацию" in result["display_text"].lower()


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
