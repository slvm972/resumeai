"""
Тесты для Cycle S1 — позиционная детекция summary/"О себе"-блока и
отдельный repositioning-промпт, применяемый ТОЛЬКО в creative-режиме.

Запуск: python -m pytest tests/test_summary_repositioning.py -v
"""
import io
import re
import sys
import pytest
from unittest.mock import patch, MagicMock

sys.path.insert(0, ".")
import missing_routes4 as mr


def _make_docx(paragraphs):
    """paragraphs: список (text, style|None)."""
    from docx import Document
    doc = Document()
    for text, style in paragraphs:
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


SUMMARY_TEXT = (
    "Experienced network engineer with a broad background in infrastructure "
    "support and troubleshooting across enterprise environments"
)
BULLET_TEXT = "Managed a team of five engineers across multiple regional offices"
PLAIN2_TEXT = "Delivered comprehensive quarterly performance reports for senior stakeholders"
HEADER_TEXT = "Experience"  # входит в SECTION_HEADERS_SET


def _docx_with_summary_before_header():
    """name(freeze) / email(freeze) / summary(plain,improve) / header(freeze) / bullet(improve)."""
    return _make_docx([
        ("John Smith", None),
        ("john.smith@example.com", None),
        (SUMMARY_TEXT, None),
        (HEADER_TEXT, None),
        (BULLET_TEXT, "List Bullet"),
    ])


def _docx_without_intro():
    """name(freeze) / email(freeze) / header(freeze, сразу) / bullet(improve) / plain(improve, ПОСЛЕ header)."""
    return _make_docx([
        ("John Smith", None),
        ("john.smith@example.com", None),
        (HEADER_TEXT, None),
        (BULLET_TEXT, "List Bullet"),
        (PLAIN2_TEXT, None),
    ])


def _fake_post_factory(captured_calls, echo_prefix="Successfully "):
    """Логирует (temperature, system_prompt, user_prompt) каждого вызова и
    возвращает echo-ответ с genuine word change (лишний префикс-слово),
    чтобы Quality Gate принимал результат без повторной попытки (retry
    не входит в объём этого цикла и не должен засорять счётчик вызовов)."""
    def _fake_post(url, headers=None, json=None, timeout=None):
        captured_calls.append({
            "temperature": json.get("temperature"),
            "system_prompt": json["messages"][0]["content"],
            "user_prompt": json["messages"][1]["content"],
        })
        user_prompt = json["messages"][1]["content"]
        out_parts = []
        for m in re.finditer(
            r"###ITEM_(\d+)###\n(.*?)(?=\n\n###ITEM_|\n\nOUTPUT|\Z)",
            user_prompt, re.DOTALL,
        ):
            iid, text = m.group(1), m.group(2).strip()
            out_parts.append(f"###ITEM_{iid}###\n{echo_prefix}{text}")
        content = "\n\n".join(out_parts)

        resp = MagicMock()
        resp.status_code = 200
        resp.json.return_value = {
            "choices": [{"message": {"content": content}}],
            "usage": {"total_tokens": 10},
        }
        return resp
    return _fake_post


# ===========================================================================
# S1_1 — precise: summary_item_id обрабатывается КАК ОБЫЧНЫЙ plain-блок,
# отдельной "summary"-группы нет вообще (поведение идентично тому, что
# было бы без Cycle S1). Docx: summary-абзац + один bullet -> ровно
# столько же HTTP-вызовов, сколько обычный bullet+plain микс (2, см.
# test_PLANB1_mixed_bullet_plain_makes_exactly_2_calls в основном
# наборе) — bullet отдельно, plain (в котором и лежит summary) отдельно.
# ===========================================================================

def test_S1_1_precise_summary_stays_plain_no_separate_group():
    docx_bytes = _docx_with_summary_before_header()
    captured = []

    with patch("requests.post", side_effect=_fake_post_factory(captured)):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="precise",
        )

    assert result["success"], result.get("error")
    assert len(captured) == 2, f"ожидали 2 вызова (bullet+plain), получили {len(captured)}"

    # Ни один вызов не должен использовать repositioning-промпт.
    for call in captured:
        assert "REPOSITION" not in call["system_prompt"], \
            "В precise-режиме summary НЕ должен получать repositioning-промпт"

    temps = sorted(c["temperature"] for c in captured)
    assert temps == sorted([0.40, 0.30]), \
        f"precise-температуры (bullet=0.40, plain=0.30) не изменились — получили {temps}"


# ===========================================================================
# S1_2 — creative: summary_item_id уходит в ОТДЕЛЬНУЮ группу со своим
# repositioning-промптом и temperature=TEMP_BY_MODE["creative"]["summary"].
# ===========================================================================

def test_S1_2_creative_summary_gets_separate_repositioning_call():
    docx_bytes = _docx_with_summary_before_header()
    captured = []

    with patch("requests.post", side_effect=_fake_post_factory(captured)):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="creative",
        )

    assert result["success"], result.get("error")
    # bullet-группа + summary-группа (plain-группа пуста — summary был
    # единственным plain-элементом и целиком ушёл в summary-группу).
    assert len(captured) == 2, f"ожидали 2 вызова (bullet+summary), получили {len(captured)}"

    summary_calls = [c for c in captured if "REPOSITION" in c["system_prompt"]]
    assert len(summary_calls) == 1, \
        f"ожидали ровно 1 вызов с repositioning-промптом, получили {len(summary_calls)}"
    assert summary_calls[0]["temperature"] == mr.TEMP_BY_MODE["creative"]["summary"]

    other_calls = [c for c in captured if "REPOSITION" not in c["system_prompt"]]
    assert len(other_calls) == 1
    assert other_calls[0]["temperature"] == mr.TEMP_BY_MODE["creative"]["bullet"]


def test_S1_2_creative_summary_item_absent_from_plain_group():
    """Regression guard на саму механику вырезания: summary НЕ должен
    попасть одновременно в plain-группу (проверяем косвенно — user_prompt
    summary-вызова содержит текст саммари, а НЕ user_prompt bullet-вызова,
    и наоборот; при дублировании один и тот же текст ушёл бы в оба)."""
    docx_bytes = _docx_with_summary_before_header()
    captured = []

    with patch("requests.post", side_effect=_fake_post_factory(captured)):
        mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="creative",
        )

    summary_call = next(c for c in captured if "REPOSITION" in c["system_prompt"])
    other_call = next(c for c in captured if "REPOSITION" not in c["system_prompt"])
    # summary-текст ушёл в summary-вызов и НЕ ушёл в bullet-вызов;
    # bullet-текст — наоборот. Если бы вырезание из group_plain_ids не
    # сработало, summary-текст оказался бы в обоих вызовах одновременно.
    assert SUMMARY_TEXT in summary_call["user_prompt"]
    assert SUMMARY_TEXT not in other_call["user_prompt"]
    assert BULLET_TEXT in other_call["user_prompt"]
    assert BULLET_TEXT not in summary_call["user_prompt"]


# ===========================================================================
# S1_3 — нет вводного абзаца (первый non-frozen item сразу section-
# заголовок) -> summary_item_id остаётся None, поведение не меняется НИ
# В КАКОМ режиме, включая creative.
# ===========================================================================

def test_S1_3_no_intro_paragraph_summary_id_stays_none_creative():
    docx_bytes = _docx_without_intro()
    captured = []

    with patch("requests.post", side_effect=_fake_post_factory(captured)):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="creative",
        )

    assert result["success"], result.get("error")
    # header(freeze) не улучшается -> 0 вызовов на него; bullet + plain
    # (идущий ПОСЛЕ header) -> 2 вызова, ни один не summary/repositioning.
    assert len(captured) == 2, f"ожидали 2 вызова (bullet+plain), получили {len(captured)}"
    for call in captured:
        assert "REPOSITION" not in call["system_prompt"], \
            "Без вводного абзаца summary НЕ должен обнаружиться даже в creative-режиме"


def test_S1_3_no_intro_paragraph_precise_unaffected():
    """То же самое в precise — не строго требуется ТЗ, но дешёвый
    regression guard на симметрию поведения."""
    docx_bytes = _docx_without_intro()
    captured = []

    with patch("requests.post", side_effect=_fake_post_factory(captured)):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="precise",
        )

    assert result["success"], result.get("error")
    assert len(captured) == 2
    for call in captured:
        assert "REPOSITION" not in call["system_prompt"]


# ===========================================================================
# Cycle S2 — фикс регрессии, найденной на реальном резюме (Дмитрий
# Соколов): "О себе" физически идёт ОТДЕЛЬНОЙ строкой-заголовком перед
# текстом summary, поэтому фиксированный i==2 из Cycle S1 промахивался
# (i==2 — сам заголовок, freeze; реальный текст summary на i==3 уходил
# в обычную plain-группу без repositioning-промпта).
#
# Новый алгоритм: пропускаем подряд идущие frozen-строки начиная с i==2,
# останавливаемся на первом же improve-элементе — если он plain, это и
# есть summary; если нет (bullet/heading/table) — отдельного summary нет,
# дальше не смотрим.
# ===========================================================================

ABOUT_ME_HEADER = "About Me"  # короткий label — freeze по длине (<15, без глагола-маркера),
                                # НЕ через SECTION_HEADERS_SET — механизм языконезависим


def _docx_header_then_summary():
    """name(freeze,i0) / email(freeze,i1) / "About Me"(freeze,i2) /
    summary-текст(plain,improve,i3) / "Experience"(freeze,i4) / bullet(improve,i5).
    Воспроизводит структуру реального резюме, на котором Cycle S1 промахнулся."""
    return _make_docx([
        ("John Smith", None),
        ("john.smith@example.com", None),
        (ABOUT_ME_HEADER, None),
        (SUMMARY_TEXT, None),
        (HEADER_TEXT, None),
        (BULLET_TEXT, "List Bullet"),
    ])


def _docx_bullet_then_plain_no_header():
    """name(freeze,i0) / email(freeze,i1) / bullet(improve,i2) /
    plain(improve,i3), БЕЗ единого заголовка — структура, идентичная той,
    что уже один раз ловила регрессию (test_CREATIVE1_mixed_bullet_plain).
    summary_item_id должен остаться None: первый non-frozen элемент
    (i==2) — bullet, не plain, поиск останавливается сразу."""
    return _make_docx([
        ("John Smith", None),
        ("john.smith@example.com", None),
        (BULLET_TEXT, "List Bullet"),
        (PLAIN2_TEXT, None),
    ])


def test_S2_1_summary_found_at_index_3_when_header_is_separate_block():
    """Заголовок "О себе"/"About Me" отдельным блоком на i==2 (freeze),
    реальный текст summary — на i==3. summary_item_id должен указывать
    именно на текст (i==3), а не на заголовок и не остаться None."""
    docx_bytes = _docx_header_then_summary()
    captured = []

    with patch("requests.post", side_effect=_fake_post_factory(captured)):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="creative",
        )

    assert result["success"], result.get("error")
    summary_calls = [c for c in captured if "REPOSITION" in c["system_prompt"]]
    assert len(summary_calls) == 1, \
        f"ожидали ровно 1 repositioning-вызов, получили {len(summary_calls)}"

    # Именно ТЕКСТ с i==3 (SUMMARY_TEXT) должен уйти в repositioning-вызов —
    # не заголовок "About Me" (он заморожен, ушёл бы токеном @@@...@@@,
    # не читаемым текстом) и не bullet.
    assert SUMMARY_TEXT in summary_calls[0]["user_prompt"]
    assert ABOUT_ME_HEADER not in summary_calls[0]["user_prompt"]
    assert BULLET_TEXT not in summary_calls[0]["user_prompt"]
    assert summary_calls[0]["temperature"] == mr.TEMP_BY_MODE["creative"]["summary"]


def test_S2_2_summary_found_at_index_2_when_no_separate_header():
    """Без отдельного заголовка — summary-текст сразу на i==2 (тот же
    сценарий, что в S1_1/S1_2, продублирован здесь под явным именем S2_2
    по требованию ТЗ Cycle S2)."""
    docx_bytes = _docx_with_summary_before_header()
    captured = []

    with patch("requests.post", side_effect=_fake_post_factory(captured)):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="creative",
        )

    assert result["success"], result.get("error")
    summary_calls = [c for c in captured if "REPOSITION" in c["system_prompt"]]
    assert len(summary_calls) == 1
    assert SUMMARY_TEXT in summary_calls[0]["user_prompt"]


def test_S2_3_regression_guard_bullet_at_index_2_gives_none():
    """Защита от регрессии: структура, идентичная test_CREATIVE1_mixed_
    bullet_plain (i==2 bullet, i==3 plain, без заголовков). Поиск должен
    остановиться на i==2 (bullet, не plain) и НЕ продолжить искать
    дальше — summary_item_id остаётся None, PLAIN2_TEXT уходит в обычную
    plain-группу без repositioning-промпта."""
    docx_bytes = _docx_bullet_then_plain_no_header()
    captured = []

    with patch("requests.post", side_effect=_fake_post_factory(captured)):
        result = mr._run_improve_pipeline(
            docx_bytes, "resume.docx", None, "fake-api-key",
            creativity_mode="creative",
        )

    assert result["success"], result.get("error")
    assert len(captured) == 2, f"ожидали 2 вызова (bullet+plain), получили {len(captured)}"
    for call in captured:
        assert "REPOSITION" not in call["system_prompt"], \
            "i==2 — bullet, не plain: поиск summary должен был остановиться здесь, " \
            "PLAIN2_TEXT на i==3 не должен был стать summary"

    temps = sorted(c["temperature"] for c in captured)
    assert temps == sorted([0.60, 0.50]), \
        f"обычные creative-температуры (bullet=0.60, plain=0.50) — получили {temps}"


# ===========================================================================
# Cycle S3 — Punctuation Sanitizer (слипшиеся предложения) + усиление
# summary-промпта (пробелы + present-tense позиционирование).
# ===========================================================================

def test_S3_1_punctuation_sanitizer_fixes_missing_space():
    result = mr._sanitize_awkward_phrasing("администрированию.Отвечал за задачи")
    assert result == "администрированию. Отвечал за задачи"


def test_S3_1_punctuation_sanitizer_handles_exclamation_and_question():
    assert mr._sanitize_awkward_phrasing("Готово!Далее следующий шаг") == \
        "Готово! Далее следующий шаг"
    assert mr._sanitize_awkward_phrasing("Уверены?Переходим к делу") == \
        "Уверены? Переходим к делу"


def test_S3_1_punctuation_sanitizer_noop_on_already_spaced_text():
    clean = "Специалист по сетям. Отвечал за инфраструктуру."
    assert mr._sanitize_awkward_phrasing(clean) == clean


def test_S3_1_regression_guard_dot_net_not_broken():
    """Ключевая находка при критической оценке ТЗ: буквально предложенный
    паттерн (?<=[.!?])(?=[А-ЯЁA-Z]) без требования к символу ПЕРЕД точкой
    ломает уже защищённые (в _PROTECT_PATTERNS/TECH_SPECIAL) технические
    термины вида "ASP.NET"/".NET" — реалистичные для IT-резюме, не
    гипотетические. Уточнённый паттерн (?<=[а-яёa-z][.!?])(?=[А-ЯЁA-Z])
    требует строчную букву перед знаком препинания — у ASP.NET перед
    точкой заглавная "P", у .NET перед точкой вообще нет буквы."""
    assert mr._sanitize_awkward_phrasing("Опыт работы с ASP.NET Framework") == \
        "Опыт работы с ASP.NET Framework"
    assert mr._sanitize_awkward_phrasing("Разработка на .NET Core") == \
        "Разработка на .NET Core"


def test_S3_1_combines_with_B2_sanitizer_in_same_text():
    """Оба санитайзера (B2 — канцелярит, S3 — пунктуация) работают в одном
    и том же тексте за один проход _sanitize_awkward_phrasing, не мешая
    друг другу."""
    result = mr._sanitize_awkward_phrasing(
        "Внёс в эксплуатацию систему.Осуществлял миграцию БД."
    )
    assert result == "Ввёл в эксплуатацию систему. Провёл миграцию БД."


def test_S3_2_summary_prompt_has_spacing_and_present_tense_rules():
    prompt = mr._build_summary_repositioning_prompt(1, "Russian")
    assert "sentence-ending punctuation" in prompt or "spacing between sentences" in prompt
    assert "PRESENT" in prompt
    assert "Отвечал за" in prompt  # bad-пример на месте
    assert "решающий" in prompt    # good-пример (причастие) на месте


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
